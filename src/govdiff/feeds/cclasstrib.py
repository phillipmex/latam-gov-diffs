"""cClassTrib - Brazilian IBS/CBS tax classification table.

Publisher: Portal Nacional da Nota Fiscal Eletronica.

The workbook has been restructured repeatedly since the first release
(2024-12-07): the sheet name changed six times, the sheet count went from one
to two, and the column set grew from 8 to 82. The parser therefore identifies
the sheet by its header content, not by name or position, and keeps every
column the publisher ships.

Every value is stored as text. The source mixes storage types across releases
(the same column is a number in one file and a string in the next), and codes
such as `000001` must not be turned into integers, so a single text
representation is the only thing that diffs cleanly across versions.
"""

from __future__ import annotations

import html
import io
import re
import unicodedata
from datetime import date, datetime

import openpyxl
import pandas as pd

from govdiff.errors import FeedError
from govdiff.fetch import fetch, make_session

FEED_ID = "cclasstrib"
BASE_URL = "https://www.nfe.fazenda.gov.br/portal/"
LISTING_URL = BASE_URL + "listaConteudo.aspx?tipoConteudo=%2FNJarYc9nus%3D"
DOCUMENT_URL = BASE_URL + "exibirArquivo.aspx?conteudo=LeNQXfyYngg="

# The listing mixes several IBS/CBS tables. Only the classification table is
# ours; the CST, cCredPres and biofuel-index tables are separate documents.
_LISTING_TITLE = re.compile(r"Classifica[cç][aã]o\s+Tribut[aá]ria\s+do\s+IBS", re.I)
_LISTING_DATE = re.compile(r"(\d{2})/(\d{2})/(\d{4})")
_LISTING_ANCHOR = re.compile(
    r"<a[^>]*href=\"(exibirArquivo\.aspx\?conteudo=[^\"]+)\"[^>]*>\s*"
    r"<span class=\"tituloConteudo\">(.*?)</span>",
    re.I | re.S,
)

# The header row of the sheet we want carries both of these.
_REQUIRED_HEADERS = ("cst-ibs/cbs", "cclasstrib")

_WHITESPACE = re.compile(r"\s+")


def _strip_tags(fragment: str) -> str:
    return html.unescape(re.sub(r"<[^>]+>", "", fragment)).strip()


def normalise_text(value: str) -> str:
    r"""Collapse every kind of whitespace the spreadsheet uses into single spaces.

    Cells mix newlines, NBSP and U+2028 line separators; Python's \s
    character class covers all of them.
    """
    return _WHITESPACE.sub(" ", value).strip()


def normalise_header(name: str) -> str:
    """`Descricao CST-IBS/CBS` -> `descricao_cst_ibs_cbs`, deterministically."""
    text = unicodedata.normalize("NFKD", normalise_text(str(name)))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^0-9A-Za-z]+", "_", text).strip("_").lower()
    return text or "column"


def normalise_value(value) -> str | None:
    """One cell -> text, or None when empty."""
    if value is None:
        return None
    if isinstance(value, datetime):
        if value.hour or value.minute or value.second:
            return value.isoformat(sep=" ")
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = normalise_text(str(value))
    return text or None


def _unique_headers(raw_headers: list) -> list[str | None]:
    """Name the columns, dropping unnamed ones and de-duplicating repeats.

    The 2025-05-06 release repeats six header labels; the 2026-06-23 release
    trails 39 unnamed columns. None marks a column to drop.
    """
    names: list[str | None] = []
    used: dict[str, int] = {}
    for cell in raw_headers:
        if cell is None or not str(cell).strip():
            names.append(None)
            continue
        name = normalise_header(cell)
        used[name] = used.get(name, 0) + 1
        if used[name] > 1:
            name = "%s_%d" % (name, used[name])
        names.append(name)
    return names


def _pick_sheet(workbook) -> tuple[object, list, int]:
    """Return (worksheet, header cells, header row index) for the cClass sheet."""
    for worksheet in workbook.worksheets:
        for row_index, row in enumerate(worksheet.iter_rows(max_row=5, values_only=True)):
            labels = {normalise_text(str(c)).lower() for c in row if c is not None}
            if all(required in labels for required in _REQUIRED_HEADERS):
                return worksheet, list(row), row_index
    raise FeedError(
        "no sheet in this workbook has both a 'CST-IBS/CBS' and a 'cClassTrib' header"
    )


# These two columns are merged vertically in the earliest releases, so only the
# first row of each CST block carries a value. Carrying the last seen value
# down is what the merged cell means.
_FORWARD_FILL = ("cst_ibs_cbs", "descricao_cst_ibs_cbs")


def parse(content: bytes) -> pd.DataFrame:
    """Parse a cClassTrib workbook into one row per classification code."""
    workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        worksheet, header_cells, header_row = _pick_sheet(workbook)
        names = _unique_headers(header_cells)
        keep = [i for i, name in enumerate(names) if name is not None]
        columns = [names[i] for i in keep]
        if "cclasstrib" not in columns:
            raise FeedError("the cClassTrib column vanished after header normalisation")
        code_position = columns.index("cclasstrib")

        rows: list[list[str | None]] = []
        carried: dict[str, str | None] = {}
        for row_index, row in enumerate(worksheet.iter_rows(values_only=True)):
            if row_index <= header_row:
                continue
            values = [normalise_value(row[i]) if i < len(row) else None for i in keep]
            if values[code_position] is None:
                # Trailing padding rows, and the blank spacer rows some
                # releases leave between CST blocks.
                continue
            for column in _FORWARD_FILL:
                if column not in columns:
                    continue
                position = columns.index(column)
                if values[position] is None:
                    values[position] = carried.get(column)
                else:
                    carried[column] = values[position]
            rows.append(values)
    finally:
        workbook.close()

    if not rows:
        raise FeedError("workbook parsed to zero classification rows")

    frame = pd.DataFrame(rows, columns=columns, dtype="string")
    return frame.reset_index(drop=True)


def list_versions(session=None, listing_html: str | None = None) -> list[dict]:
    """Every dated cClassTrib release the portal indexes, oldest first.

    Fetches the listing page exactly once per call. Pass `listing_html` to
    parse an already-downloaded page and make no request at all.
    """
    if listing_html is None:
        session = session or make_session()
        result = fetch(LISTING_URL, session=session, conditional=False)
        listing_html = (result.content or b"").decode("utf-8", errors="replace")

    entries: list[dict] = []
    for match in _LISTING_ANCHOR.finditer(listing_html):
        title = _strip_tags(match.group(2))
        if not _LISTING_TITLE.search(title):
            continue
        stamp = _LISTING_DATE.search(title)
        if not stamp:
            continue
        day, month, year = stamp.groups()
        href = html.unescape(match.group(1))
        # ASP.NET puts a raw base64 token in the query string; a literal '+'
        # arrives here decoded as a space and has to go back.
        href = href.replace(" ", "+")
        entries.append(
            {
                "title": title,
                "date": "%s-%s-%s" % (year, month, day),
                "published": "%s/%s/%s" % (day, month, year),
                "url": BASE_URL + href,
            }
        )

    # The portal lists the same release only once, but sort defensively.
    seen: set[str] = set()
    unique = []
    for entry in sorted(entries, key=lambda e: (e["date"], e["url"])):
        if entry["url"] in seen:
            continue
        seen.add(entry["url"])
        unique.append(entry)
    return unique


def current_document(session=None, listing_html: str | None = None) -> dict:
    """The newest release on the listing page.

    `DOCUMENT_URL` is only the release that was current when this feed was
    registered. Each release has its own permanent URL, so the nightly run has
    to re-read the index to notice a new one.
    """
    entries = list_versions(session=session, listing_html=listing_html)
    if not entries:
        raise FeedError("the listing page shows no dated cClassTrib release")
    return entries[-1]
